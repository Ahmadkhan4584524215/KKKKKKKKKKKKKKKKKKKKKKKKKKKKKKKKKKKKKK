import io
import json
import os
import re
import time

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document

APP_TITLE = "Resume ATS Analyzer"
MODEL_NAME = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
FALLBACK_MODEL_NAME = os.getenv("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash")
MAX_RETRIES = 3
RETRY_DELAY_SECONDS = 3
MAX_RESUME_CHARS = 50000

st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")


def get_api_key():
    """Read the Gemini API key from Streamlit secrets or an environment variable."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY")


def extract_text(uploaded_file):
    """Extract readable text from TXT, PDF, or DOCX resumes."""
    filename = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if filename.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages).strip()

    if filename.endswith(".docx"):
        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs).strip()

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


def clean_json_text(text):
    """Convert a Gemini response into a Python dictionary, tolerating code fences."""
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
        if not match:
            raise ValueError("Gemini returned an invalid JSON response.")
        return json.loads(match.group(0))


def analyze_resume(resume_text, job_description, api_key):
    """Ask Gemini to score and improve the resume."""
    client = genai.Client(api_key=api_key)

    job_text = job_description.strip() if job_description else ""
    if not job_text:
        job_text = "No job description was provided. Give a general ATS-readiness assessment instead of pretending to match a specific job."

    prompt = f"""
You are an expert resume reviewer and ATS (Applicant Tracking System) optimization assistant.
Analyze the resume below.

IMPORTANT RULES:
1. Do not invent jobs, education, skills, metrics, certifications, dates, or achievements.
2. The ATS score is an estimated readiness score, not a score from a real employer's ATS.
3. If a job description is provided, evaluate keyword and requirement alignment against it.
4. If no job description is provided, give a general ATS-readiness score.
5. Penalize hard-to-parse formatting only when the extracted text gives evidence of it. Be cautious because text extraction cannot fully reveal visual formatting.
6. Recommendations must be practical and specific.
7. Keep the response valid JSON only. No Markdown and no code fences.

Return exactly this JSON structure:
{{
  "ats_score": 0,
  "score_label": "Poor | Needs Work | Good | Strong",
  "summary": "Short overall assessment",
  "section_scores": {{
    "keyword_match": 0,
    "experience": 0,
    "skills": 0,
    "clarity": 0,
    "ats_format_readiness": 0
  }},
  "strengths": ["..."],
  "improvements": ["..."],
  "missing_keywords": ["..."],
  "ats_warnings": ["..."],
  "action_plan": ["..."],
  "rewritten_bullets": [
    {{"original": "Existing bullet from resume", "suggestion": "Improved version without inventing facts"}}
  ]
}}

JOB DESCRIPTION:
{job_text[:30000]}

RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""

    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )

    # Retry temporary Gemini 503/429 errors. If the primary model is busy,
    # fall back to another Flash model.
    models_to_try = [MODEL_NAME]
    if FALLBACK_MODEL_NAME != MODEL_NAME:
        models_to_try.append(FALLBACK_MODEL_NAME)

    last_error = None

    for model_name in models_to_try:
        for attempt in range(MAX_RETRIES):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                    config=config,
                )
                return clean_json_text(response.text)

            except Exception as exc:
                last_error = exc
                error_text = str(exc).upper()

                # Retry only temporary capacity/rate-limit failures.
                if "503" not in error_text and "UNAVAILABLE" not in error_text and "429" not in error_text:
                    raise

                if attempt < MAX_RETRIES - 1:
                    time.sleep(RETRY_DELAY_SECONDS * (attempt + 1))

        # Primary model may be temporarily overloaded; try fallback.

    raise RuntimeError(
        "Gemini is temporarily unavailable after multiple retries. "
        f"Primary model: {MODEL_NAME}. Fallback: {FALLBACK_MODEL_NAME}. "
        "Please try again in a minute."
    ) from last_error


def show_score(score):
    score = max(0, min(100, int(score)))
    st.metric("Estimated ATS Score", f"{score}/100")
    st.progress(score / 100)


def safe_list(value):
    return value if isinstance(value, list) else []


st.title("📄 Resume ATS Analyzer")
st.caption("Upload a resume and get an AI-powered ATS-readiness score plus practical improvements.")

with st.sidebar:
    st.header("How it works")
    st.write("1. Upload your resume")
    st.write("2. Optionally paste a job description")
    st.write("3. Gemini analyzes the resume")
    st.write("4. Review the score and improvements")
    st.divider()
    st.caption(f"Primary Gemini model: `{MODEL_NAME}`")
    st.caption(f"Fallback model: `{FALLBACK_MODEL_NAME}`")

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Supported formats: PDF, DOCX, TXT. For best results, use a text-based PDF rather than a scanned image PDF.",
)

job_description = st.text_area(
    "Job description (optional)",
    height=220,
    placeholder="Paste the job description here for a more targeted ATS score...",
)

analyze_clicked = st.button("🔍 Analyze Resume", type="primary", use_container_width=True)

if analyze_clicked:
    if uploaded_file is None:
        st.error("Please upload a resume first.")
        st.stop()

    api_key = get_api_key()
    if not api_key:
        st.error("Gemini API key not found. Add GEMINI_API_KEY to your local Streamlit secrets or environment variables.")
        st.stop()

    try:
        with st.spinner("Reading your resume and analyzing it with Gemini..."):
            resume_text = extract_text(uploaded_file)
            if not resume_text:
                st.error("No readable text was found. If this is a scanned PDF, use a text-based PDF or DOCX/TXT file.")
                st.stop()

            result = analyze_resume(resume_text, job_description, api_key)

        score = int(result.get("ats_score", 0))
        st.success("Analysis complete!")
        show_score(score)

        col1, col2 = st.columns(2)
        with col1:
            st.subheader("Overall assessment")
            st.write(result.get("summary", "No summary returned."))
            st.write(f"**Rating:** {result.get('score_label', 'N/A')}")
        with col2:
            st.subheader("Score breakdown")
            section_scores = result.get("section_scores", {})
            for key, value in section_scores.items():
                label = key.replace("_", " ").title()
                try:
                    numeric_value = max(0, min(100, int(value)))
                except (TypeError, ValueError):
                    numeric_value = 0
                st.write(f"**{label}:** {numeric_value}/100")
                st.progress(numeric_value / 100)

        st.divider()
        left, right = st.columns(2)
        with left:
            st.subheader("✅ Strengths")
            for item in safe_list(result.get("strengths")):
                st.write(f"- {item}")

            st.subheader("🔑 Missing / weak keywords")
            for item in safe_list(result.get("missing_keywords")):
                st.write(f"- {item}")

        with right:
            st.subheader("⚠️ ATS warnings")
            for item in safe_list(result.get("ats_warnings")):
                st.write(f"- {item}")

            st.subheader("🛠 Improvements")
            for item in safe_list(result.get("improvements")):
                st.write(f"- {item}")

        st.subheader("🎯 Action plan")
        for index, item in enumerate(safe_list(result.get("action_plan")), start=1):
            st.write(f"**{index}.** {item}")

        rewritten = safe_list(result.get("rewritten_bullets"))
        if rewritten:
            st.subheader("✍️ Suggested bullet improvements")
            for item in rewritten:
                if isinstance(item, dict):
                    st.markdown(f"**Original:** {item.get('original', '')}")
                    st.markdown(f"**Suggestion:** {item.get('suggestion', '')}")
                    st.divider()

        st.download_button(
            "⬇️ Download analysis as JSON",
            data=json.dumps(result, indent=2, ensure_ascii=False),
            file_name="resume_ats_analysis.json",
            mime="application/json",
        )

    except Exception as exc:
        error_text = str(exc)
        if "temporarily unavailable" in error_text.lower() or "503" in error_text or "UNAVAILABLE" in error_text:
            st.error("Gemini is temporarily overloaded. The app retried automatically, but both models are currently unavailable.")
            st.info("Please wait about a minute and click Analyze Resume again.")
        else:
            st.error(f"Analysis failed: {exc}")
            st.info("Check your API key, internet connection, supported file type, and Gemini model availability.")

st.divider()
st.caption("Note: This tool provides an AI estimate for resume improvement. It is not a guarantee of how a particular employer's ATS will score a resume.")
